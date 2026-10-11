"""
Generates FICTIONAL sample data for testing Sift (no real people).
  python3 scripts/make_samples.py
Requires: pip install reportlab python-docx
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

OUT = Path(__file__).resolve().parent.parent / "sample-data"
OUT.mkdir(exist_ok=True)

JD = """AI/ML Engineer – LLM Applications
Company: Northwind Analytics (fictional)
Location: Bengaluru, India (Hybrid – 3 days in office)
Experience: 3–6 years
Employment type: Full-time

About the role
We are building an AI assistant that answers questions over enterprise documents. You will design, build and ship LLM-powered features end to end, from retrieval pipelines to production APIs.

What you will do
- Build retrieval-augmented generation (RAG) pipelines over large document collections
- Design prompts, evaluation sets and guardrails for LLM features
- Develop and maintain Python APIs (FastAPI) that serve models in production
- Work with vector databases (Pinecone, FAISS or similar) and embeddings
- Deploy and monitor services on AWS using Docker
- Collaborate with product and data teams to measure quality and cost

Must have
- 3+ years of professional experience in Python
- Hands-on experience with LLMs (OpenAI, Claude, Groq or open-source models)
- RAG and LangChain or LlamaIndex
- Vector databases and embeddings
- FastAPI or Flask for REST APIs
- Docker and AWS
- SQL

Nice to have
- Fine-tuning (LoRA / PEFT), Hugging Face Transformers
- LLM evaluation frameworks (RAGAS, DeepEval)
- Kubernetes, CI/CD

Education
B.E./B.Tech/M.Tech in Computer Science or a related field."""

CVS = {
    "01_Ananya_Rao.pdf": """Ananya Rao
Bengaluru | ananya.rao@example.com | +91 98450 11111

SUMMARY
Machine learning engineer with 5 years of Python experience, the last 3 building LLM products in production.

EXPERIENCE
Senior ML Engineer, Lumen Docs (fictional) — Jan 2022 to Present
- Built a RAG pipeline over 2 million enterprise PDFs using LangChain, OpenAI embeddings and Pinecone; cut answer latency by 40%.
- Designed prompt templates and an evaluation set of 600 questions; tracked faithfulness with RAGAS.
- Shipped FastAPI services in Docker on AWS (ECS, S3, Lambda) serving 50k queries a day.
- Fine-tuned a Llama 3 8B model with LoRA (Hugging Face PEFT) for domain classification.

Data Engineer, Brightline Insights (fictional) — Jul 2020 to Dec 2021
- Wrote Python ETL jobs and complex SQL on PostgreSQL; set up CI/CD with GitHub Actions.

EDUCATION
B.Tech, Computer Science, 2020

SKILLS
Python, SQL, LangChain, Pinecone, FAISS, FastAPI, Docker, AWS, RAGAS, Hugging Face, LoRA""",

    "02_Rahul_Verma.docx": """Rahul Verma
Pune | rahul.verma@example.com | +91 98220 22222

PROFILE
ML engineer with 3.5 years building recommendation and NLP models in Python.

WORK EXPERIENCE
Machine Learning Engineer, Kestrel Retail (fictional) — Apr 2022 to Present
- Trained PyTorch recommendation models; served them through Flask REST APIs in Docker on Google Cloud (GCP).
- Built a 6-month prototype support chatbot with LlamaIndex, FAISS and an open-source LLM (Mistral 7B).
- Wrote SQL feature pipelines on BigQuery.

Data Analyst, Kestrel Retail (fictional) — Jan 2021 to Mar 2022
- Reporting in SQL and Python (pandas).

EDUCATION
M.Tech, Data Science, 2020

SKILLS
Python, PyTorch, Flask, Docker, GCP, SQL, LlamaIndex, FAISS, Hugging Face Transformers""",

    "03_Priya_Nair.pdf": """Priya Nair
Kochi | priya.nair@example.com

EDUCATION
B.E., Computer Science, 2025

EXPERIENCE
Associate Software Engineer, Tidewater Labs (fictional) — Jul 2025 to Present
- Maintain internal Python tools and SQL reports.

ML Intern, Tidewater Labs (fictional) — Jan 2025 to Jun 2025
- Built a RAG chatbot over HR policies using LangChain, Chroma and the OpenAI API, served with FastAPI.

PROJECTS
- Resume parser with spaCy and Python.
- Dockerised a FastAPI demo app.

SKILLS
Python, SQL, LangChain, Chroma, FastAPI, Docker (basic)""",

    "04_Karthik_Subramanian.docx": """Karthik Subramanian
Chennai | karthik.s@example.com | +91 94440 44444

SUMMARY
Backend engineer with 5 years in Java microservices.

EXPERIENCE
Senior Software Engineer, Harbor Payments (fictional) — 2021 to Present
- Java 17 and Spring Boot microservices on AWS (EKS, RDS); Docker and Kubernetes; Kafka event streaming.
- Optimised SQL queries on MySQL, reducing p95 latency by 30%.

Software Engineer, Harbor Payments (fictional) — 2020 to 2021
- REST APIs in Spring Boot; Jenkins CI/CD.

EDUCATION
B.E., Information Technology, 2020

SKILLS
Java, Spring Boot, AWS, Docker, Kubernetes, Kafka, SQL, Jenkins, some Python scripting""",

    "05_Meera_Iyer.pdf": """Meera Iyer
Hyderabad | meera.iyer@example.com

SKILLS
Python, LLM, RAG, LangChain, LlamaIndex, Pinecone, FAISS, Vector Databases, FastAPI, Flask, Docker, AWS, Kubernetes, SQL, Hugging Face, LoRA, RAGAS, Prompt Engineering

EXPERIENCE
Customer Support Executive, Skyline Telecom (fictional) — 2022 to Present
- Resolve 60+ customer tickets a day; maintain ticket notes in the CRM.

Data Entry Operator, Velvet Logistics (fictional) — 2020 to 2022
- Entered shipment records into spreadsheets with 99% accuracy.

CERTIFICATIONS
- Online course: Introduction to Generative AI (completed 2025)

EDUCATION
B.Com, 2020""",

    "06_Arjun_Mehta.docx": """Arjun Mehta
Mumbai | arjun.mehta@example.com | +91 98200 66666

PERSONAL DETAILS
Gender: Male
Age: 31
Date of Birth: 14/03/1995
Marital Status: Married
Nationality: Indian
Religion: Hindu
Photo: attached

SUMMARY
Data scientist with 6 years in NLP for fintech, recently focused on generative AI.

EXPERIENCE
Lead Data Scientist, Orbit Fintech (fictional) — 2023 to Present
- Built a RAG prototype for policy Q&A with LangChain, OpenAI and pgvector; ran prompt experiments.
- Deployed BERT classifiers with Hugging Face on AWS SageMaker, packaged in Docker.

Data Scientist, Orbit Fintech (fictional) — 2020 to 2023
- NLP models for KYC document classification in Python; SQL on Redshift.

EDUCATION
M.Tech, Computer Science, 2020

SKILLS
Python, NLP, Hugging Face, LangChain, pgvector, AWS SageMaker, Docker, SQL""",

    "07_Sneha_Kulkarni.pdf": """Sneha Kulkarni
Bengaluru | sneha.k@example.com

SUMMARY
Full-stack developer with 4 years building React and Node.js web apps.

EXPERIENCE
Software Engineer II, Maple Commerce (fictional) — 2022 to Present
- Built React and Node.js (Express) features for a checkout platform.
- Added an AI product-description generator calling the OpenAI API from Node.js.
- Deployed services with Docker on AWS (EC2, S3).

Software Engineer, Maple Commerce (fictional) — 2021 to 2022
- REST APIs in Node.js with PostgreSQL (SQL).

EDUCATION
B.E., Computer Science, 2021

SKILLS
JavaScript, TypeScript, React, Node.js, SQL, Docker, AWS, some Python""",
}


def write_pdf(path: Path, text: str):
    styles = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=styles["Normal"], fontSize=10, leading=14)
    head = ParagraphStyle("h", parent=styles["Heading2"], fontSize=12, spaceBefore=8)
    title = ParagraphStyle("t", parent=styles["Title"], fontSize=18, alignment=0)
    story = []
    for i, line in enumerate(text.split("\n")):
        line = line.strip().replace("&", "&amp;")
        if not line:
            story.append(Spacer(1, 4))
        elif i == 0:
            story.append(Paragraph(line, title))
        elif line.isupper():
            story.append(Paragraph(line, head))
        else:
            story.append(Paragraph(line, body))
    SimpleDocTemplate(str(path), pagesize=A4, topMargin=40, bottomMargin=40).build(story)


def write_docx(path: Path, text: str):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    for i, line in enumerate(text.split("\n")):
        line = line.strip()
        if not line:
            continue
        if i == 0:
            doc.add_heading(line, level=0)
        elif line.isupper():
            doc.add_heading(line.title(), level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)
    doc.save(str(path))


if __name__ == "__main__":
    write_docx(OUT / "JD_AI_ML_Engineer.docx", JD)
    write_pdf(OUT / "JD_AI_ML_Engineer.pdf", JD)
    cv_dir = OUT / "cvs"
    cv_dir.mkdir(exist_ok=True)
    for name, text in CVS.items():
        (write_pdf if name.endswith(".pdf") else write_docx)(cv_dir / name, text)
    print("Wrote sample data to", OUT)
